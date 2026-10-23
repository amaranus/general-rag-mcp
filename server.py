# -*- coding: utf-8 -*-
"""General local file RAG MCP server for Unsloth Desktop.

Vektör araması + SQLite FTS5 (BM25) hibrit arama, başlık-farkında chunking.
"""
from __future__ import annotations
import hashlib, json, os, re, sqlite3, sys
from pathlib import Path
from typing import Any
import numpy as np
from mcp.server.fastmcp import FastMCP
from sentence_transformers import SentenceTransformer

def load_env_file(path: Path) -> None:
    """APP_DIR/.env dosyasını okur; mevcut ortam değişkenlerini ezmez."""
    if not path.is_file(): return
    for raw in path.read_text(encoding='utf-8-sig', errors='replace').splitlines():
        line=raw.strip()
        if not line or line.startswith('#'): continue
        if line.lower().startswith('export '): line=line[7:].strip()
        if '=' not in line: continue
        key,value=line.split('=',1); key=key.strip(); value=value.strip()
        if value[:1] in '\'"':
            q=value[0]; end=value.find(q,1)
            if end!=-1: value=value[1:end]
        else: value=re.split(r'\s+#',value,maxsplit=1)[0].strip()
        if key: os.environ.setdefault(key,value)

APP_DIR = Path(__file__).resolve().parent
load_env_file(Path(os.environ.get('RAG_ENV_FILE', str(APP_DIR / '.env'))).expanduser())
WORKSPACE = Path(os.environ.get('RAG_WORKSPACE', str(APP_DIR / 'workspace'))).expanduser().resolve()
DB_PATH = Path(os.environ.get('RAG_DB_PATH', str(APP_DIR / 'rag_index.sqlite3'))).expanduser().resolve()
MODEL_NAME = os.environ.get('RAG_EMBEDDING_MODEL', 'sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2')
CHUNK_SIZE, CHUNK_OVERLAP = int(os.getenv('RAG_CHUNK_SIZE', '550')), int(os.getenv('RAG_CHUNK_OVERLAP', '100'))
MAX_FILE_MB = int(os.getenv('RAG_MAX_FILE_MB', '100'))
RRF_K = int(os.getenv('RAG_RRF_K', '60'))
HEADING_BOOST = float(os.getenv('RAG_HEADING_BOOST', '0.45'))
mcp = FastMCP('General File RAG')
_model = None
TEXT_EXTS = {'.txt','.md','.markdown','.rst','.csv','.tsv','.json','.jsonl','.xml','.html','.htm','.rtf','.log','.yaml','.yml','.ini','.toml','.py','.js','.ts','.jsx','.tsx','.css','.sql','.sh','.ps1','.bat','.c','.h','.cpp','.hpp','.java','.go','.rs','.php','.rb','.ipynb'}
SUPPORTED = TEXT_EXTS | {'.pdf','.docx','.xlsx','.pptx','.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp'}

CHAPTER_RE = re.compile(r'^\s*(chapter|bölüm|bolum|kısım|kisim|part|section|appendix|ek)\b[\s.:\-–]*([IVXLCDM]+|\d+)?\b.*$', re.I)
NUM_RE = re.compile(r'^\s*[IVXLCDM]{2,7}[.)]?\s*$')
NUMHEAD_RE = re.compile(r'^\s*\d{1,2}(?:\.\d{1,2})*[.)]?\s+[A-ZÇĞİÖŞÜ]')
MD_RE = re.compile(r'^\s{0,3}#{1,6}\s+\S')
ALLCAPS_RE = re.compile(r"^[A-ZÇĞİÖŞÜ][A-ZÇĞİÖŞÜ0-9 .,'’\-:!?()]{2,}$")
PAGE_RE = re.compile(r'^\[(?:sayfa|page)\s+\d+\]$', re.I)
TOKEN_RE = re.compile(r'\w+', re.UNICODE)

SCHEMA_SQL = """CREATE TABLE IF NOT EXISTS chunks (
 id INTEGER PRIMARY KEY,
 path TEXT NOT NULL,
 file_hash TEXT NOT NULL,
 chunk_no INTEGER NOT NULL,
 heading TEXT NOT NULL DEFAULT '',
 content TEXT NOT NULL,
 embedding BLOB NOT NULL,
 embedding_dim INTEGER NOT NULL,
 UNIQUE(path,file_hash,chunk_no));
CREATE INDEX IF NOT EXISTS idx_chunks_path ON chunks(path);"""

FTS_SQL = """CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(content, heading, content='chunks', content_rowid='id', tokenize='unicode61 remove_diacritics 2');
CREATE TRIGGER IF NOT EXISTS chunks_ai AFTER INSERT ON chunks BEGIN INSERT INTO chunks_fts(rowid,content,heading) VALUES(new.id,new.content,new.heading); END;
CREATE TRIGGER IF NOT EXISTS chunks_ad AFTER DELETE ON chunks BEGIN INSERT INTO chunks_fts(chunks_fts,rowid,content,heading) VALUES('delete',old.id,old.content,old.heading); END;
CREATE TRIGGER IF NOT EXISTS chunks_au AFTER UPDATE ON chunks BEGIN INSERT INTO chunks_fts(chunks_fts,rowid,content,heading) VALUES('delete',old.id,old.content,old.heading); INSERT INTO chunks_fts(rowid,content,heading) VALUES(new.id,new.content,new.heading); END;"""

def _fts_supported() -> bool:
    try:
        t=sqlite3.connect(':memory:'); t.execute('CREATE VIRTUAL TABLE t USING fts5(x)'); t.close(); return True
    except Exception: return False

HAS_FTS = _fts_supported()

def ensure_workspace(): WORKSPACE.mkdir(parents=True, exist_ok=True)
def safe_path(value: str) -> Path:
    ensure_workspace(); p = Path(value).expanduser()
    if not p.is_absolute(): p = WORKSPACE / p
    p = p.resolve()
    try: p.relative_to(WORKSPACE)
    except ValueError: raise ValueError(f'Yol workspace dışında. Dosyayı buraya koy: {WORKSPACE}')
    if not p.exists(): raise FileNotFoundError(f'Bulunamadı: {p}')
    return p

def db():
    c=sqlite3.connect(DB_PATH, timeout=30); c.row_factory=sqlite3.Row
    cols=[r['name'] for r in c.execute('PRAGMA table_info(chunks)').fetchall()]
    if cols and 'heading' not in cols:
        c.executescript('DROP TABLE IF EXISTS chunks; DROP TABLE IF EXISTS chunks_fts;')
        print('Şema güncellendi: eski indeks temizlendi, dosyaları yeniden indeksle.', file=sys.stderr, flush=True)
    c.executescript(SCHEMA_SQL)
    if HAS_FTS: c.executescript(FTS_SQL)
    c.commit(); return c

def model():
    global _model
    if _model is None:
        print(f'Embedding modeli yükleniyor: {MODEL_NAME}', file=sys.stderr, flush=True)
        _model=SentenceTransformer(MODEL_NAME)
        print('Embedding modeli hazır.', file=sys.stderr, flush=True)
    return _model

def extract_text(p: Path) -> str:
    ext=p.suffix.lower()
    if ext in TEXT_EXTS:
        raw=p.read_text(encoding='utf-8-sig', errors='replace')
        if ext in {'.html','.htm'}:
            from bs4 import BeautifulSoup
            return BeautifulSoup(raw,'html.parser').get_text('\n',strip=True)
        if ext=='.rtf':
            from striprtf.striprtf import rtf_to_text
            return rtf_to_text(raw)
        if ext=='.ipynb':
            try:
                nb=json.loads(raw); return '\n\n'.join(''.join(c.get('source',[])) if isinstance(c.get('source',[]),list) else str(c.get('source','')) for c in nb.get('cells',[]))
            except Exception: return raw
        return raw
    if ext=='.pdf':
        import pymupdf
        with pymupdf.open(p) as doc:
            return '\n\n'.join(f'[Sayfa {i}]\n{page.get_text("text")}' for i,page in enumerate(doc,1) if page.get_text('text').strip())
    if ext=='.docx':
        from docx import Document
        d=Document(p); parts=[x.text for x in d.paragraphs if x.text.strip()]
        for t in d.tables:
            for r in t.rows: parts.append(' | '.join(c.text.strip() for c in r.cells))
        return '\n'.join(parts)
    if ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(p,read_only=True,data_only=True); out=[]
        try:
            for ws in wb.worksheets:
                out.append(f'[Sayfa: {ws.title}]')
                for row in ws.iter_rows(values_only=True):
                    vals=[str(v) for v in row if v is not None]
                    if vals: out.append(' | '.join(vals))
        finally: wb.close()
        return '\n'.join(out)
    if ext=='.pptx':
        from pptx import Presentation
        prs=Presentation(p); out=[]
        for i,s in enumerate(prs.slides,1):
            vals=[sh.text for sh in s.shapes if getattr(sh,'has_text_frame',False) and sh.text.strip()]
            if vals: out.append(f'[Slayt {i}]\n'+'\n'.join(vals))
        return '\n\n'.join(out)
    if ext in {'.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp'}:
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(p),lang=os.getenv('RAG_OCR_LANG','tur+eng'))
    raise ValueError(f'Desteklenmeyen dosya türü: {ext}')

def is_heading(line: str) -> bool:
    s=line.strip()
    if not s or len(s)>100 or PAGE_RE.match(s): return False
    if MD_RE.match(s) or NUMHEAD_RE.match(s) or NUM_RE.match(s): return True
    if CHAPTER_RE.match(s) and not s.endswith('.'): return True
    if ALLCAPS_RE.match(s) and len(re.findall(r'[A-ZÇĞİÖŞÜ]',s))>=3: return True
    return False

def split_sections(text: str) -> list[tuple[str,str]]:
    text=text.replace('\r\n','\n').replace('\r','\n'); lines=text.split('\n')
    sections: list[tuple[str,str]]=[]; heading=''; body: list[str]=[]
    def flush():
        nonlocal heading, body
        if body and '\n'.join(body).strip(): sections.append((heading,'\n'.join(body)))
        heading=''; body=[]
    i=0
    while i<len(lines):
        line=lines[i]; s=line.strip()
        if is_heading(line):
            flush(); title=re.sub(r'^#{1,6}\s*','',s)
            if CHAPTER_RE.match(s) and not s.endswith('.') and i+1<len(lines):
                nxt=lines[i+1].strip()
                if nxt and len(nxt)<=100 and is_heading(lines[i+1]) and not CHAPTER_RE.match(nxt) and not nxt.endswith('.'):
                    title=f'{title} — {re.sub(r"^#{1,6}\s*","",nxt)}'; i+=1
            heading=title
        else: body.append(line)
        i+=1
    flush()
    return sections or [('',text)]

def chunks(text: str) -> list[str]:
    text=re.sub(r'\n{3,}','\n\n',text.replace('\x00',' ')).strip(); out=[]; start=0
    while start<len(text):
        end=min(start+CHUNK_SIZE,len(text))
        if end<len(text):
            b=max(text.rfind('\n\n',start,end),text.rfind('\n',start,end),text.rfind(' ',start,end))
            if b>start+int(CHUNK_SIZE*.55): end=b
        s=text[start:end].strip()
        if s: out.append(s)
        if end>=len(text): break
        start=max(end-CHUNK_OVERLAP,start+1)
    return out

def file_hash(p: Path) -> str:
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()

def index_one(p: Path, force=False) -> dict[str,Any]:
    if not p.is_file(): raise ValueError(f'Dosya değil: {p}')
    if p.suffix.lower() not in SUPPORTED: raise ValueError(f'Desteklenmeyen uzantı: {p.suffix}')
    size=p.stat().st_size/(1024*1024)
    if size>MAX_FILE_MB: raise ValueError(f'Dosya {size:.1f} MB; sınır {MAX_FILE_MB} MB')
    rel=p.relative_to(WORKSPACE).as_posix(); digest=file_hash(p); c=db()
    try:
        old=c.execute('SELECT COUNT(*) n FROM chunks WHERE path=? AND file_hash=?',(rel,digest)).fetchone()['n']
        if old and not force: return {'path':rel,'status':'already_indexed','chunks':old}
        text=extract_text(p)
        parts: list[tuple[str,str]]=[]
        for heading,body in split_sections(text):
            for part in chunks(body): parts.append((heading,part))
        if not parts: return {'path':rel,'status':'no_text_extracted','chunks':0,'message':'Metin çıkarılamadı. Taranmış PDF için OCR gerekir.'}
        inputs=[f'{h}\n{t}' if h else t for h,t in parts]
        vectors=model().encode(inputs,convert_to_numpy=True,normalize_embeddings=True,show_progress_bar=False).astype(np.float32)
        c.execute('DELETE FROM chunks WHERE path=?',(rel,))
        for i,((heading,part),v) in enumerate(zip(parts,vectors)):
            c.execute('INSERT INTO chunks(path,file_hash,chunk_no,heading,content,embedding,embedding_dim) VALUES(?,?,?,?,?,?,?)',(rel,digest,i,heading,part,v.tobytes(),int(v.size)))
        c.commit()
        return {'path':rel,'status':'indexed','chunks':len(parts),'sections':len({h for h,_ in parts if h}),'characters_extracted':len(text),'embedding_model':MODEL_NAME}
    finally: c.close()

def _query_tokens(query: str) -> list[str]:
    seen=set(); out=[]
    for t in (x.lower() for x in TOKEN_RE.findall(query)):
        if t and t not in seen: seen.add(t); out.append(t)
    return out[:32]

def norm(s: str) -> str:
    return re.sub(r'[^\w]+',' ',s.lower()).strip()

def _fts_ranked_ids(c: sqlite3.Connection, query: str, path_filter: str, limit: int) -> list[int]:
    if not HAS_FTS: return []
    toks=_query_tokens(query)
    if not toks: return []
    match=' OR '.join(f'"{t}"'+('*' if len(t)>2 else '') for t in toks)
    sql='SELECT chunks_fts.rowid AS id FROM chunks_fts WHERE chunks_fts MATCH ?'
    params: list[Any]=[match]
    if path_filter.strip():
        sql+=' AND chunks_fts.rowid IN (SELECT id FROM chunks WHERE path LIKE ?)'; params.append('%'+path_filter.strip()+'%')
    sql+=' ORDER BY bm25(chunks_fts,1.0,5.0) LIMIT ?'; params.append(max(limit,50))
    try: return [r['id'] for r in c.execute(sql,params).fetchall()]
    except sqlite3.OperationalError: return []

def _vector_scores(c: sqlite3.Connection, query: str, path_filter: str, rows: list[sqlite3.Row]) -> dict[int,float]:
    if not rows: return {}
    q=model().encode([query],convert_to_numpy=True,normalize_embeddings=True,show_progress_bar=False)[0].astype(np.float32)
    sims: dict[int,float]={}
    for r in rows:
        v=np.frombuffer(r['embedding'],dtype=np.float32)
        if v.size==q.size: sims[r['id']]=float(q@v)
    return sims

def _neighbor(c: sqlite3.Connection, path: str, chunk_no: int) -> str:
    r=c.execute('SELECT content FROM chunks WHERE path=? AND chunk_no=?',(path,chunk_no)).fetchone()
    return r['content'] if r else ''

def _resolve_heading(c: sqlite3.Connection, rel: str, heading: str) -> str | None:
    wanted=norm(heading)
    if not wanted: return None
    titles=[r['heading'] for r in c.execute('SELECT DISTINCT heading FROM chunks WHERE path=? AND heading<>\'\'',(rel,)).fetchall()]
    best=None; best_score=0.0
    for cand in titles:
        nc=norm(cand)
        if wanted==nc: return cand
        qs,cs=set(wanted.split()),set(nc.split())
        score=2.0 if (wanted in nc or nc in wanted) else (len(qs&cs)/len(qs) if qs else 0.0)
        if score>best_score: best_score, best = score, cand
    return best if best_score>=0.5 else None

@mcp.tool()
def list_files() -> str:
    """Workspace içindeki desteklenen dosyaları listeler."""
    ensure_workspace(); items=[]
    for p in sorted(WORKSPACE.rglob('*')):
        if p.is_file() and p.suffix.lower() in SUPPORTED: items.append({'path':p.relative_to(WORKSPACE).as_posix(),'size_bytes':p.stat().st_size,'extension':p.suffix.lower()})
    return json.dumps({'workspace':str(WORKSPACE),'files':items},ensure_ascii=False,indent=2)

@mcp.tool()
def index_file(path: str, force: bool=False) -> str:
    """Workspace içindeki tek dosyayı metne dönüştürüp yerel RAG indeksine ekler. path workspace'e göreli olabilir."""
    try: return json.dumps(index_one(safe_path(path),force),ensure_ascii=False,indent=2)
    except Exception as e: return json.dumps({'error':str(e)},ensure_ascii=False)

@mcp.tool()
def index_directory(path: str='.', recursive: bool=True, force: bool=False) -> str:
    """Bir klasördeki desteklenen dosyaları indeksler."""
    try:
        target=safe_path(path)
        if not target.is_dir(): return json.dumps({'error':'Yol klasör değil.'},ensure_ascii=False)
        it=target.rglob('*') if recursive else target.iterdir(); results=[]
        for p in sorted(it):
            if p.is_file() and p.suffix.lower() in SUPPORTED:
                try: results.append(index_one(p,force))
                except Exception as e: results.append({'path':str(p),'status':'error','error':str(e)})
        return json.dumps({'processed':len(results),'results':results},ensure_ascii=False,indent=2)
    except Exception as e: return json.dumps({'error':str(e)},ensure_ascii=False)

@mcp.tool()
def search_documents(query: str, limit: int=5, path_filter: str='', expand_neighbors: bool=False) -> str:
    """İndekslenmiş dosyalarda hibrit (vektör + BM25) arama yapar; başlık eşleşmelerini ve yakın chunk'ları öne çıkarır."""
    if not str(query).strip(): return 'Aranacak soru veya konu gir.'
    limit=max(1,min(int(limit),20)); c=db()
    try:
        sql='SELECT id,path,chunk_no,heading,content,embedding FROM chunks'
        params: list[Any]=[]
        if path_filter.strip(): sql+=' WHERE path LIKE ?'; params.append('%'+path_filter.strip()+'%')
        rows=c.execute(sql,params).fetchall()
        if not rows: return json.dumps({'query':query,'results':[],'message':'İndeks boş. Önce index_file veya index_directory kullan.'},ensure_ascii=False,indent=2)
        sims=_vector_scores(c,query,path_filter,rows)
        vec_order=[cid for cid,_ in sorted(sims.items(),key=lambda kv:kv[1],reverse=True)]
        fts_ids=_fts_ranked_ids(c,query,path_filter,limit)
        fused: dict[int,float]={}
        for rank,cid in enumerate(vec_order): fused[cid]=fused.get(cid,0.0)+1.0/(RRF_K+rank+1)
        for rank,cid in enumerate(fts_ids): fused[cid]=fused.get(cid,0.0)+1.0/(RRF_K+rank+1)
        by_id={r['id']:r for r in rows}; qtoks=[t for t in _query_tokens(query) if len(t)>1]
        nq=norm(query)
        for cid,score in list(fused.items()):
            r=by_id.get(cid)
            if r is None: continue
            h=norm(r['heading'])
            if h and qtoks:
                hits=sum(1 for t in qtoks if t in h)
                if hits: score*=1.0+HEADING_BOOST*hits/len(qtoks)
                if nq and nq in h: score*=1.25
            fused[cid]=score
        ranked=sorted(fused.items(),key=lambda kv:kv[1],reverse=True)[:limit]
        results=[]
        for cid,score in ranked:
            r=by_id[cid]; item={'path':r['path'],'chunk':r['chunk_no']+1,'heading':r['heading'],'similarity':round(sims.get(cid,0.0),4),'hybrid_score':round(score,5),'text':r['content']}
            if expand_neighbors:
                item['context_before']=_neighbor(c,r['path'],r['chunk_no']-1)[-600:]
                item['context_after']=_neighbor(c,r['path'],r['chunk_no']+1)[:600]
            results.append(item)
        return json.dumps({'query':query,'count':len(results),'search_mode':'hybrid (vector+bm25)' if HAS_FTS else 'vector','results':results},ensure_ascii=False,indent=2)
    except Exception as e: return json.dumps({'error':str(e)},ensure_ascii=False)
    finally: c.close()

@mcp.tool()
def list_sections(path: str='') -> str:
    """İndekslenmiş dosyaların bölüm/başlık listesini verir; bölüm adından arama için kullanılabilir."""
    try:
        c=db()
        try:
            sql='SELECT path,heading,COUNT(*) chunks,MIN(chunk_no)+1 first_chunk,MAX(chunk_no)+1 last_chunk FROM chunks'
            params: list[Any]=[]
            if path.strip():
                p=safe_path(path); rel=p.relative_to(WORKSPACE).as_posix(); sql+=' WHERE path=?'; params.append(rel)
            sql+=' GROUP BY path,heading ORDER BY path,first_chunk'
            return json.dumps({'sections':[dict(r) for r in c.execute(sql,params).fetchall()]},ensure_ascii=False,indent=2)
        finally: c.close()
    except Exception as e: return json.dumps({'error':str(e)},ensure_ascii=False)

@mcp.tool()
def read_section(path: str, heading: str='', max_chars: int=30000) -> str:
    """İndeksteki bir bölümün tüm chunk'larını sırayla döndürür. heading boşsa dosyanın tamamını verir."""
    try:
        p=safe_path(path); rel=p.relative_to(WORKSPACE).as_posix(); c=db()
        try:
            if heading.strip():
                title=_resolve_heading(c,rel,heading)
                if not title:
                    available=[r['heading'] for r in c.execute('SELECT DISTINCT heading FROM chunks WHERE path=? AND heading<>\'\' LIMIT 20',(rel,)).fetchall()]
                    return json.dumps({'error':f'Bölüm bulunamadı: {heading}','available_headings':available},ensure_ascii=False)
                rows=c.execute('SELECT chunk_no,heading,content FROM chunks WHERE path=? AND heading=? ORDER BY chunk_no',(rel,title)).fetchall()
            else: rows=c.execute('SELECT chunk_no,heading,content FROM chunks WHERE path=? ORDER BY chunk_no',(rel,)).fetchall()
            if not rows: return json.dumps({'error':f'Bölüm bulunamadı: {heading}'},ensure_ascii=False)
            text='\n\n'.join(r['content'] for r in rows); n=max(1000,min(int(max_chars),60000))
            return json.dumps({'path':rel,'heading':rows[0]['heading'],'matched_headings':sorted({r['heading'] for r in rows if r['heading']}),'chunks':len(rows),'characters':len(text),'text':text[:n],'truncated':len(text)>n},ensure_ascii=False,indent=2)
        finally: c.close()
    except Exception as e: return json.dumps({'error':str(e)},ensure_ascii=False)

@mcp.tool()
def read_file(path: str, max_chars: int=12000) -> str:
    """Workspace içindeki dosyadan metin çıkarır; indeksleme yapmaz."""
    try:
        p=safe_path(path)
        if p.suffix.lower() not in SUPPORTED: return f'Desteklenmeyen dosya türü: {p.suffix}'
        text=extract_text(p); n=max(500,min(int(max_chars),50000))
        return json.dumps({'path':p.relative_to(WORKSPACE).as_posix(),'characters':len(text),'text':text[:n],'truncated':len(text)>n},ensure_ascii=False,indent=2)
    except Exception as e: return json.dumps({'error':str(e)},ensure_ascii=False)

@mcp.tool()
def remove_from_index(path: str) -> str:
    """Bir dosyanın indeksini siler; asıl dosyayı silmez."""
    try:
        p=safe_path(path); rel=p.relative_to(WORKSPACE).as_posix(); c=db()
        try:
            cur=c.execute('DELETE FROM chunks WHERE path=?',(rel,)); c.commit()
            return json.dumps({'path':rel,'deleted_chunks':cur.rowcount},ensure_ascii=False)
        finally: c.close()
    except Exception as e: return json.dumps({'error':str(e)},ensure_ascii=False)

if __name__=='__main__':
    ensure_workspace(); print(f'RAG workspace: {WORKSPACE}',file=sys.stderr,flush=True); mcp.run(transport='stdio')
